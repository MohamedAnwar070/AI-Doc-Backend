from __future__ import annotations

from docx import Document

from core.word_worker import WordWorker


def test_word_worker_extracts_ordered_text_and_layout_html(tmp_path) -> None:
    docx_path = tmp_path / "notice.docx"
    document = Document()
    document.add_heading("Notice", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Important ").bold = True
    paragraph.add_run("message")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Section"
    table.cell(1, 1).text = "42"
    document.save(docx_path)

    worker = WordWorker(
        job_id="job-1",
        docx_path=docx_path,
        storage=None,
        progress_tracker=None,
        resume_handler=None,
    )

    text, elements = worker._extract_docx()

    assert "Notice" in text
    assert "Important message" in text
    assert "Name\tValue" in text
    assert [el["type"] for el in elements] == ["heading", "paragraph", "table"]
    assert "<strong>Important </strong>" in elements[1]["html"]
    assert elements[2]["html"].startswith('<table class="docx-table">')
